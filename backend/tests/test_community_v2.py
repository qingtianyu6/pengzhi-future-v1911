from datetime import UTC, datetime, timedelta
import io
import json
import pytest
from sqlalchemy import select
from app.database import SessionLocal, initialize_database
from app.models import User
from app.models.community import ForumPost, NewsArticle, NewsDigest, AutomationJob, AITrace, KnowledgeDocument
from app.services import news_service, content_network, forum_service
from test_final_product import create_context

def post(client, path, payload=None, expected=200):
    response = client.post('/api/community/' + path, json=payload or {})
    assert response.status_code == expected, response.text
    return response.json()['data']

def create_post(client, **kw):
    return post(client, 'forum/posts', {'title': '番茄清晨叶面湿润观察', 'body': '记录了作物阶段和现场湿度，想知道下一步应补充哪些观察。', 'category': 'question', **kw}, 201)

def register(client, phone='13900000001'):
    result = client.post('/api/auth/register', json={'display_name': '另一位棚友', 'phone': phone, 'organization': '独立组织', 'password': 'Greenhouse2026'})
    assert result.status_code == 200
    return {'Authorization': 'Bearer ' + result.json()['data']['token']}

def test_forum_full_flow(client):
    gid, _ = create_context(client, 'FORUM')
    p = create_post(client, fields={'crop': 'tomato'}, tags=['湿度'])
    p = post(client, f"forum/posts/{p['id']}/comments", {'body': '建议记录同一区域处理前后情况'}, 201)
    first = p['comments'][0]['id']
    p = post(client, f"forum/posts/{p['id']}/comments", {'body': '补充记录了区域差异', 'parent_id': first}, 201)
    assert p['comment_count'] == 2 and p['comments'][1]['parent_id'] == first
    response = client.put(f"/api/community/forum/posts/{p['id']}/accepted/{first}")
    assert response.status_code == 200 and response.json()['data']['accepted_comment_id'] == first
    for enabled in (True, True, False):
        response = client.put(f"/api/community/forum/posts/{p['id']}/reactions/bookmark", json={'enabled': enabled})
        assert response.status_code == 200 and response.json()['data']['bookmarked'] == enabled
    post(client, f"forum/posts/{p['id']}/reports", {'reason': '请核实经验适用范围'})
    reports = client.get('/api/community/forum/moderation/reports').json()['data']
    assert len(reports) == 1
    post(client, f"forum/posts/{p['id']}/moderation", {'action': 'review', 'note': '已审核为讨论参考，仍需实际观察验证'})
    assert client.get('/api/community/forum/moderation/reports').json()['data'][0]['status'] == 'resolved'
    analysis = post(client, f"forum/posts/{p['id']}/assist", {'use_model': False})
    assert analysis['published'] is False and analysis['mode'] == 'knowledge_retrieval'
    r = post(client, f"forum/posts/{p['id']}/research", {'greenhouse_id': gid})
    assert r['research_id']
    post(client, f"forum/posts/{p['id']}/research", {'greenhouse_id': gid, 'research_id': r['research_id'], 'confirmed': True})
    edited = client.get(f"/api/community/forum/posts/{p['id']}").json()['data']
    edit = client.put(f"/api/community/forum/posts/{p['id']}", json={**edited, 'body': '修改后的观察问题内容。', 'asset_ids': []})
    assert edit.status_code == 200 and edit.json()['data']['review_state'] == 'unreviewed'
    stale = client.put(f"/api/community/forum/posts/{p['id']}", json={**edited, 'body': '旧版本覆盖请求'})
    assert stale.status_code == 409

def test_visibility_permissions_and_filters(client):
    private = create_post(client, visibility='private')
    org = create_post(client, visibility='organization')
    public = create_post(client, visibility='public', tags=['测试'], fields={'crop': 'tomato'})
    headers = register(client)
    for item in (private, org):
        assert client.get(f"/api/community/forum/posts/{item['id']}", headers=headers).status_code == 404
    assert client.get(f"/api/community/forum/posts/{public['id']}", headers=headers).status_code == 200
    assert client.delete(f"/api/community/forum/posts/{public['id']}", headers=headers).status_code == 403
    assert client.post(f"/api/community/forum/posts/{public['id']}/moderation", json={'action': 'hide', 'note': '无权审核'}, headers=headers).status_code == 403
    listing = client.get('/api/community/forum/posts', params={'crop': 'tomato', 'tag': '测试'}, headers=headers).json()['data']
    assert listing['total'] == 1
    p2 = post(client, 'forum/posts', {'title': '另一个问题', 'body': '独立的讨论内容'}, 201)
    c = post(client, f"forum/posts/{p2['id']}/comments", {'body': '这是一条回答'}, 201)['comments'][0]
    assert client.post(f"/api/community/forum/posts/{public['id']}/comments", json={'body': '错误跨帖回复', 'parent_id': c['id']}).status_code == 422
    assert client.get('/api/community/forum/posts').status_code == 200

def test_attachments_and_markdown_filter(client, monkeypatch, tmp_path):
    monkeypatch.setattr(forum_service, 'ASSET_DIR', tmp_path)
    assert client.post('/api/community/forum/posts', json={'title': '恶意内容', 'body': '<script>alert(1)</script>'}).status_code == 422
    invalid = client.post('/api/community/forum/assets', files={'file': ('bad.png', b'not-an-image', 'image/png')})
    assert invalid.status_code == 422
    from PIL import Image
    content = io.BytesIO(); Image.new('RGB', (30, 30), 'green').save(content, format='PNG')
    uploaded = client.post('/api/community/forum/assets', files={'file': ('leaf.png', content.getvalue(), 'image/png')})
    assert uploaded.status_code == 201
    aid = uploaded.json()['data']['id']
    private = create_post(client, visibility='private', asset_ids=[aid])
    assert client.get(f'/api/community/forum/assets/{aid}').status_code == 200
    listed = client.get('/api/community/forum/posts').json()['data']['items']
    assert next(p for p in listed if p['id'] == private['id'])['assets'] == [{'id': aid, 'filename': 'leaf.png', 'media_type': uploaded.json()['data']['media_type']}]
    headers = register(client)
    assert client.get(f'/api/community/forum/assets/{aid}', headers=headers).status_code == 404
    assert private['id'] not in [p['id'] for p in client.get('/api/community/forum/posts', headers=headers).json()['data']['items']]
    assert client.post('/api/community/forum/posts', json={'title': '偷用附件', 'body': '不能引用他人的附件', 'asset_ids': [aid]}, headers=headers).status_code == 422
    assert client.get('/api/community/forum/assets/999999').status_code == 404

def rss():
    date = datetime.now(UTC).strftime('%a, %d %b %Y %H:%M:%S GMT')
    return f'<?xml version="1.0"?><rss version="2.0"><channel><title>测试农业机构</title><link>https://example.org/</link><description>Fixture</description><item><title>番茄种植管理与天气观察</title><link>https://example.org/agriculture/1?utm_source=rss</link><pubDate>{date}</pubDate></item><item><title>番茄种植管理与天气观察</title><link>https://example.org/agriculture/1</link><pubDate>{date}</pubDate></item></channel></rss>'.encode()

def add_source(client, **kw):
    return post(client, 'news/sources', {'name': '测试农业机构', 'url': 'https://example.org/rss.xml',
        'kind': 'rss', 'permission_confirmed': True, 'permission_note': '自动化测试来源，不代表实际机构授权', **kw}, 201)

def test_rss_digest_idempotency_and_preferences(client, monkeypatch):
    monkeypatch.setattr(content_network, 'fetch', lambda *a, **kw: (rss(), {'etag': 'fixture'}, 'https://example.org/rss.xml', 200))
    s = add_source(client)
    result = post(client, f"news/sources/{s['id']}/refresh")
    assert result['added'] == 1
    again = post(client, f"news/sources/{s['id']}/refresh")
    assert again['status'] == 'already_processed'
    a = client.get('/api/community/news/articles').json()['data']['items'][0]
    assert a['url'] == 'https://example.org/agriculture/1' and a['status'] == 'metadata_only'
    assert client.put('/api/community/news/preferences', json={'crops': ['tomato'], 'topics': [], 'enabled': True, 'local_hour': 8}).status_code == 200
    first = post(client, 'news/digests')
    second = post(client, 'news/digests')
    assert first['digest_id'] == second['digest_id']
    digests = client.get('/api/community/news/digests').json()['data']
    assert len(digests) == 1 and len(digests[0]['articles']) == 1
    assert client.put(f"/api/community/news/digests/{first['digest_id']}/read").status_code == 200
    assert client.put(f"/api/community/news/bookmarks/{a['id']}", json={'bookmarked': True, 'read': True}).json()['data']['bookmarked']
    assert client.get('/api/community/news/articles', params={'bookmarks': True}).json()['data']['total'] == 1
    with SessionLocal() as db:
        assert news_service.claim(db, 'exclusive-test') is not None
        assert news_service.claim(db, 'exclusive-test') is None
        job = db.get(AutomationJob, 'exclusive-test'); job.lease_until = datetime.now(UTC).replace(tzinfo=None)-timedelta(seconds=1); db.commit()
        assert news_service.claim(db, 'exclusive-test') is not None
    headers = register(client)
    assert client.get('/api/community/news/articles', headers=headers).json()['data']['total'] == 0
    assert client.post(f"/api/community/news/sources/{s['id']}/refresh", headers=headers).status_code == 404

def test_collection_failure_retry_and_empty_digest(client, monkeypatch):
    def fail(*a, **kw): raise OSError('fixture network outage')
    monkeypatch.setattr(content_network, 'fetch', fail)
    s = add_source(client)
    result = post(client, f"news/sources/{s['id']}/refresh")
    assert result['status'] == 'failed'
    source = client.get('/api/community/news/sources').json()['data'][0]
    assert source['failures'] == 1 and source['next_retry_at']
    monkeypatch.setattr(content_network, 'fetch', lambda *a, **kw: (rss(), {}, 'https://example.org/rss.xml', 200))
    retry = post(client, f"news/sources/{s['id']}/refresh")
    assert retry['added'] == 1
    assert client.put('/api/community/news/preferences', json={'crops': ['muskmelon'], 'topics': [], 'enabled': True, 'local_hour': 8}).status_code == 200
    assert post(client, 'news/digests')['status'] == 'no_new_content'
    assert client.get('/api/community/news/digests').json()['data'] == []

@pytest.mark.parametrize('url', ['http://127.0.0.1/a', 'http://[::1]/a', 'http://169.254.169.254/a', 'file:///etc/passwd', 'https://localhost/rss', 'https://user:password@example.org/rss', 'http://example.org:22/a'])
def test_network_local_targets_denied(url):
    from app.services.errors import ServiceError
    with pytest.raises(ServiceError): content_network.public_url(url, resolve=False)

def test_document_import_search_review_and_isolation(client):
    text = '# 番茄叶面观察\n\n番茄管理需要记录叶面湿润比例、生育阶段和观察区域。比较处理前后的观察指标，并核实天气和灌溉等伴随变化。'
    response = client.post('/api/community/knowledge/documents', files={'file': ('guide.md', text.encode(), 'text/markdown')},
        data={'title': '番茄观察指导', 'crops': 'tomato', 'source_url': 'https://example.org/guide', 'rights_confirmed': 'true'})
    assert response.status_code == 201, response.text
    d = response.json()['data']
    assert d['status'] == 'ready' and d['chunk_count'] >= 1
    assert all(not i['id'].startswith('document-') for i in client.get('/api/community/knowledge/search', params={'q': '叶面湿润比例', 'crop': 'tomato'}).json()['data']['items'])
    assert client.put(f"/api/community/knowledge/documents/{d['id']}/review", json={'approved': True}).status_code == 200
    found = client.get('/api/community/knowledge/search', params={'q': '番茄叶面湿润比例', 'crop': 'tomato'}).json()['data']['items']
    assert any(i['id'].startswith('document-') and i['position'] for i in found)
    chunks = {chunk['id']: chunk for chunk in client.get(f"/api/community/knowledge/documents/{d['id']}").json()['data']['chunks']}
    for item in found:
        for source in item['sources']:
            if source.get('document_id') == d['id']:
                assert source['chunk_id'] in chunks
                assert source['position'] == chunks[source['chunk_id']]['position']
    melon = client.get('/api/community/knowledge/search', params={'q': '番茄叶面湿润比例', 'crop': 'muskmelon'}).json()['data']['items']
    assert not any(i['id'].startswith('document-') for i in melon)
    assert client.get('/api/community/knowledge/search', params={'q': 'qzxv987654'}).json()['data']['items'] == []
    headers = register(client)
    assert client.get(f"/api/community/knowledge/documents/{d['id']}", headers=headers).status_code == 404
    assert client.get('/api/community/knowledge/documents', headers=headers).json()['data'] == []
    assert client.get('/api/community/knowledge/search', params={'q': '番茄观察指导'}, headers=headers).status_code == 200

def test_docx_and_failed_pdf_are_real_parse_results(client):
    from docx import Document
    d = Document(); d.add_heading('设施番茄观察', 0); d.add_paragraph('这是用于自动化测试的农业文档，记录番茄种植阶段、区域、观察时间和指标。比较观察时需要保留来源，不把前后变化直接归因。')
    blob = io.BytesIO(); d.save(blob)
    response = client.post('/api/community/knowledge/documents', files={'file': ('book.docx', blob.getvalue())}, data={'rights_confirmed': 'true', 'crops': 'tomato'})
    assert response.status_code == 201 and response.json()['data']['parser'] == 'python-docx'
    broken = client.post('/api/community/knowledge/documents', files={'file': ('scan.pdf', b'broken-pdf')}, data={'rights_confirmed': 'true'})
    assert broken.status_code == 201 and broken.json()['data']['status'] == 'failed' and broken.json()['data']['error']

def test_local_evaluation_and_tracing_contract(client):
    result = post(client, 'ai/evaluations', {'mode': 'local'}, 201)
    assert result['total'] == 8 and result['mode'] == 'local_contract'
    assert len(client.get('/api/community/ai/evaluations').json()['data']) == 1
    assert client.get('/api/community/ai/traces').json()['data']['actual_cost'] is None
    initialize_database(); initialize_database()
    assert client.get('/api/community/forum/posts').status_code == 200


def test_model_evaluation_requires_consent_and_records_actual_provider_output(client, monkeypatch):
    from app.services import ai_provider
    calls = []
    class Provider:
        last_usage = {'total_tokens': 12}
        async def chat(self, messages, tools):
            calls.append(messages)
            context = json.loads(messages[-1]['content'])['context']
            return {'content': json.dumps({'answer': '需要补充观察并人工复核', 'sources': [r['id'] for r in context.get('references', [])],
                'crop': context.get('crop', ''), 'uncertainty': '缺少现场证据', 'requires_confirmation': True,
                'requires_manual_review': True, 'actions': []})}
    monkeypatch.setattr(ai_provider, 'get_llm_provider', lambda: Provider())
    assert client.post('/api/community/ai/evaluations', json={'mode': 'model'}).status_code == 422
    assert calls == []
    result = post(client, 'ai/evaluations', {'mode': 'model', 'paid_model_authorized': True}, 201)
    assert len(calls) == 8 and result['total'] == 8 and result['passed'] == 8
    assert all(r['evidence']['usage']['total_tokens'] == 12 for r in result['results'])
    other = register(client)
    assert client.get('/api/community/ai/evaluations', headers=other).json()['data'] == []


def test_worker_restart_dedup_and_link_status(client, monkeypatch):
    import asyncio
    calls = []
    def fetch(*args, **kwargs):
        calls.append(args)
        return rss(), {}, 'https://example.org/rss.xml', 200
    monkeypatch.setattr(content_network, 'fetch', fetch)
    add_source(client)
    client.put('/api/community/news/preferences', json={'crops': ['tomato'], 'topics': [], 'enabled': True, 'local_hour': 0})
    asyncio.run(news_service.tick())
    asyncio.run(news_service.tick())
    assert len(calls) == 1
    assert len(client.get('/api/community/news/digests').json()['data']) == 1
    a = client.get('/api/community/news/articles').json()['data']['items'][0]
    assert post(client, f"news/articles/{a['id']}/check-link")['available']
    def gone(*args, **kwargs): raise OSError('fixture unavailable')
    monkeypatch.setattr(content_network, 'fetch', gone)
    assert not post(client, f"news/articles/{a['id']}/check-link")['available']
    stored = client.get('/api/community/news/articles').json()['data']['items'][0]
    assert stored['analysis']['link_status'] == 'unavailable'


def test_news_shutdown_waits_for_inflight_collection(client, monkeypatch):
    import asyncio
    from threading import Event
    started, release, finished = Event(), Event(), Event()
    add_source(client)
    def controlled_collection(_source_id):
        started.set()
        assert release.wait(5)
        finished.set()
    monkeypatch.setattr(news_service, 'collect_source', controlled_collection)
    async def scenario():
        task = asyncio.create_task(news_service.tick())
        assert await asyncio.to_thread(started.wait, 5)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.is_set()
    asyncio.run(scenario())


@pytest.mark.parametrize('kind', ['rss', 'web'])
def test_atom_and_authorized_web_use_real_parsers(client, monkeypatch, kind):
    if kind == 'rss':
        raw = b'<feed xmlns="http://www.w3.org/2005/Atom"><title>Farm feed</title><entry><title>Tomato cultivation observation</title><link href="https://example.org/atom-article"/><updated>2026-10-16T00:00:00Z</updated><id>test-one</id></entry></feed>'
    else:
        paragraph = 'Tomato cultivation requires recording observation time and crop stages. This synthetic article is only a parser fixture. '
        raw = ('<html><head><title>Tomato cultivation observation</title></head><body><article><h1>Tomato cultivation observation</h1>' + ''.join('<p>'+paragraph+'</p>' for _ in range(10)) + '</article></body></html>').encode()
    monkeypatch.setattr(content_network, 'fetch', lambda *a, **kw: (raw, {}, 'https://example.org/article', 200))
    source = add_source(client, kind=kind)
    result = post(client, f"news/sources/{source['id']}/refresh")
    assert result['status'] == 'completed' and result['added'] == 1
    stored = client.get('/api/community/news/articles').json()['data']['items'][0]
    assert stored['title'] == 'Tomato cultivation observation'
    assert stored['status'] == 'metadata_only' and 'fixture' not in stored['summary']


def test_pdf_text_and_page_position(client):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 40 700 Td (Tomato observation test: record date, crop stage and greenhouse area before comparing changes.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    blob = io.BytesIO(); writer.write(blob)
    response = client.post('/api/community/knowledge/documents', files={'file': ('guide.pdf', blob.getvalue())}, data={'rights_confirmed': 'true', 'crops': 'tomato'})
    assert response.status_code == 201, response.text
    doc = response.json()['data']
    assert doc['status'] == 'ready' and doc['parser'] == 'pypdf'
    detail = client.get(f"/api/community/knowledge/documents/{doc['id']}").json()['data']
    assert 'Tomato observation' in detail['chunks'][0]['content']
    assert '1' in detail['chunks'][0]['position']
