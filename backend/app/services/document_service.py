"""Small local document pipeline with page/paragraph locators and explicit failures."""
from datetime import UTC, datetime
import hashlib
import io
import re
from pathlib import Path
import zipfile
from fastapi.encoders import jsonable_encoder
from sqlalchemy import delete, select
from app.models.community import KnowledgeChunk, KnowledgeDocument
from app.services.access import actor, require_role
from app.services.content_network import public_url
from app.services.errors import ServiceError

def document(db, document_id):
    d = db.scalar(select(KnowledgeDocument).where(KnowledgeDocument.id == document_id,
        KnowledgeDocument.organization_id == actor().organization_id))
    if d is None: raise ServiceError(404, "知识文档不存在")
    return d

def data(db, d, detail=False):
    chunks = list(db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == d.id).order_by(KnowledgeChunk.id)))
    result = {k: getattr(d, k) for k in ("id", "title", "filename", "source_url", "source_quality", "crops", "stages", "applicability", "status", "parser", "error", "created_at", "updated_at")}
    result["chunk_count"] = len(chunks)
    result["can_delete"] = d.owner_id == actor().id or actor().role == "admin"
    if detail: result["chunks"] = [{"id": c.id, "position": c.position, "content": c.content} for c in chunks]
    return jsonable_encoder(result)

def parse(content, suffix):
    if suffix in (".md", ".txt"):
        return "utf8-markdown", [("段落 " + str(i+1), p.strip()) for i, p in enumerate(re.split(r"\n\s*\n", content.decode("utf-8"))) if p.strip()]
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted: raise ValueError("加密PDF无法解析")
        if len(reader.pages) > 600: raise ValueError("单次最多600页，请按章节拆分")
        return "pypdf", [(f"第{i+1}页", page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if suffix == ".docx":
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist()) > 40_000_000: raise ValueError("DOCX展开体积过大")
            if any(x.filename.endswith("vbaProject.bin") for x in z.infolist()): raise ValueError("不支持宏文档")
        d = Document(io.BytesIO(content))
        blocks = [(f"段落{i+1}", p.text) for i, p in enumerate(d.paragraphs)]
        for i, table in enumerate(d.tables):
            blocks.append((f"表格{i+1}", "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)))
        return "python-docx", blocks
    raise ValueError("仅支持PDF、DOCX、Markdown和UTF-8文本")

def ingest(db, filename, content, title, crops, stages, source_url, applicability, rights_confirmed):
    require_role("admin", "agronomist")
    if not rights_confirmed: raise ServiceError(422, "请确认有权导入并用于内部检索")
    if not content or len(content) > 15_000_000: raise ServiceError(413, "文档为空或超过15MB")
    if source_url: public_url(source_url, resolve=False)
    digest = hashlib.sha256(content).hexdigest()
    existing = db.scalar(select(KnowledgeDocument).where(KnowledgeDocument.organization_id == actor().organization_id, KnowledgeDocument.sha256 == digest))
    if existing: return data(db, existing, True)
    d = KnowledgeDocument(organization_id=actor().organization_id, owner_id=actor().id,
        title=title.strip()[:200] or Path(filename).stem[:200], filename=Path(filename).name[:180],
        sha256=digest, crops=crops, stages=stages, source_url=source_url or None, applicability=applicability[:3000])
    db.add(d)
    try: db.flush()
    except Exception:
        db.rollback()
        duplicate = db.scalar(select(KnowledgeDocument).where(KnowledgeDocument.organization_id == actor().organization_id, KnowledgeDocument.sha256 == digest))
        if duplicate: return data(db, duplicate, True)
        raise
    try:
        parser, blocks = parse(content, Path(filename).suffix.lower())
        chunks = []
        for position, text in blocks:
            text = text.replace("\x00", "").strip()
            if not text: continue
            for i in range(0, len(text), 1000):
                piece = text[i:i+1000].strip()
                if piece: chunks.append(KnowledgeChunk(document_id=d.id, position=f"{position} · 片段{i//1000+1}", content=piece))
        if sum(len(c.content) for c in chunks) < 40: raise ValueError("有效文本不足40字；扫描PDF需要OCR或复杂解析")
        if len(chunks) > 6000: raise ValueError("内容过长，请分章节导入")
        db.add_all(chunks); d.parser = parser; d.status = "ready"
    except Exception as exc:
        d.status = "failed"; d.error = str(exc)[:300]; d.parser = "failed"
    db.commit(); db.refresh(d)
    return data(db, d, True)

def review(db, document_id, approved):
    require_role("admin")
    d = document(db, document_id)
    if d.status != "ready": raise ServiceError(422, "文档尚未成功解析")
    d.source_quality = "reviewed" if approved else "unverified"
    d.updated_at = datetime.now(UTC); db.commit()
    return data(db, d)

def remove(db, document_id):
    d = document(db, document_id)
    if d.owner_id != actor().id and actor().role != "admin": raise ServiceError(403, "无权删除此文档")
    db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == d.id)); db.delete(d); db.commit()
    return {"deleted": True}

def docling_probe():
    import importlib.util
    return {"installed": importlib.util.find_spec("docling") is not None, "mode": "optional_experimental",
        "note": "基础导入使用本地轻量解析；Docling需要独立环境与模型资源，不自动下载权重，不将未运行的复杂解析标为成功。"}
